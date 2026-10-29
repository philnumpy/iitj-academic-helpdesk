"""Normalise the timetable spreadsheets into tidy CSV files used by the schedule/email tools.

* courses.csv              - every course floated in AY 2026-27 Semester I
* slot_timings.csv         - slot -> (day, time) from the two slot-system grids
* minor_exam_schedule.csv  - slot -> (date, time) of the minor examination
* major_exam_schedule.csv  - slot -> (date, time) of the major examination
"""
import json
import re
from pathlib import Path

import openpyxl
import pandas as pd

from .config import COURSES_CSV, EXAMS_CSV, MAJOR_EXAMS_CSV, RAW_DIR, SLOTS_CSV

TT_DIR = RAW_DIR / "timetables"
UG2_FILE = TT_DIR / "AY 2026-27_Semester I_Slot System (2nd Year onwards).xlsx"
UG1_FILE = TT_DIR / "AY 2026-27 Semester-I Slot System (B.Tech. 1st Year).xlsx"
FINAL_FILE = TT_DIR / "Finalized courses list.xlsx"

DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri"]

# Sheet names of "Finalized courses list.xlsx" -> department names used in Classroom Allocation
SHEET_DEPT = {
    "EE": "Electrical Engineering", "ECE": "Electronics Engineering",
    "SoLA": "School of Liberal Arts (SoLA)", "CET": "Center for Education Technology (CET)",
    "CIE": "Civil and Infrastructure Engineering", "MT": "Materials Engineering",
    "Aero": "Aerospace Engineering", "EnAuto": "Energy and Automotive Engineering",
    "BSBE": "Bioscience and Bioengineering", "Math": "Mathematics",
    "Chemical": "Chemical Engineering", "AIDE": "School of Artificial Intelligence and Data Science (AIDE)",
    "CSE": "Computer Science and Engineering", "MedTech": "Medical Technology Center",
    "SME": "School of Management and Entrepreneurship (SME)", "Physics": "Physics",
    "Chemistry": "Chemistry", "SOD": "School of Design", "ME": "Mechanical Engineering",
    "DESE": "Environmental and Sustainability Engineering",
}

COLUMNS = ["code", "name", "ltp", "credits", "slot", "department", "instructor",
           "co_instructor", "classroom", "remarks", "source"]


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _norm_code(code: str) -> str:
    return re.sub(r"\s+", "", code).upper()


# --------------------------------------------------------------------------- courses
def _classroom_allocation() -> list[dict]:
    ws = openpyxl.load_workbook(UG2_FILE, data_only=True)["Classroom Allocation"]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        code, name = _s(r[0]), _s(r[1])
        if not code and not name:
            continue
        out.append(dict(code=_norm_code(code), name=name, ltp=_s(r[2]), credits=_s(r[3]),
                        slot=_s(r[4]), department=_s(r[5]), instructor=_s(r[6]),
                        co_instructor=_s(r[7]), classroom=_s(r[9]), remarks="",
                        source=f"{UG2_FILE.name} (sheet: Classroom Allocation)"))
    return out


def _ug1_courses() -> list[dict]:
    ws = openpyxl.load_workbook(UG1_FILE, data_only=True)["Courses as per Curriculum"]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        code, name = _s(r[1]), _s(r[2])
        if not code and not name:
            continue
        out.append(dict(code=_norm_code(code), name=name, ltp=_s(r[5]), credits=_s(r[4]),
                        slot=_s(r[8]), department="B.Tech. 1st Year (common)",
                        instructor=_s(r[6]), co_instructor="", classroom=_s(r[9]),
                        remarks=f"UG 1st year, batch {_s(r[3])}" + (f"; email {_s(r[7])}" if r[7] else ""),
                        source=f"{UG1_FILE.name} (sheet: Courses as per Curriculum)"))
    return out


def _map_header(cells: list[str]) -> dict:
    """Map heterogeneous column headers of the department sheets onto COLUMNS."""
    m = {}
    for i, h in enumerate(c.lower() for c in cells):
        if not h:
            continue
        if "co-instructor" in h:
            m.setdefault("co_instructor", i)
        elif any(k in h for k in ("instructor", "faculty", "coordinator")):
            m.setdefault("instructor", i)
        elif "code" in h:
            m.setdefault("code", i)
        elif "lab slot" in h:
            m.setdefault("remarks_lab", i)
        elif "slot" in h:
            m.setdefault("slot", i)
        elif "l-t-p" in h or h == "ltp" or "[l-t-p]" in h:
            m.setdefault("ltp", i)
        elif "credit" in h:
            m.setdefault("credits", i)
        elif any(k in h for k in ("name", "title")):
            m.setdefault("name", i)
        elif any(k in h for k in ("remark", "comment", "offered", "category", "type", "program", "year", "batch")):
            m.setdefault("remarks_" + str(i), i)
    return m


def _finalized_list() -> list[dict]:
    wb = openpyxl.load_workbook(FINAL_FILE, data_only=True)
    out = []
    for ws in wb.worksheets:
        header, section = None, ""
        for r in ws.iter_rows(values_only=True):
            cells = [_s(c) for c in r]
            if not any(cells):
                continue
            if any("code" in c.lower() for c in cells) and any("name" in c.lower() or "title" in c.lower() for c in cells):
                header = _map_header(cells)
                continue
            if header is None or sum(bool(c) for c in cells) == 1:
                section = next(c for c in cells if c)  # e.g. "PG Semester-1 (DADT)"
                continue
            get = lambda k: cells[header[k]] if k in header and header[k] < len(cells) else ""
            name = get("name")
            if not name:
                continue
            remarks = "; ".join(cells[i] for k, i in header.items() if k.startswith("remarks") and cells[i])
            if section:
                remarks = f"{section}; {remarks}" if remarks else section
            out.append(dict(code=_norm_code(get("code")), name=name, ltp=get("ltp"), credits=get("credits"),
                            slot=get("slot"), department=SHEET_DEPT.get(ws.title, ws.title),
                            instructor=get("instructor"), co_instructor=get("co_instructor"),
                            classroom="", remarks=remarks,
                            source=f"{FINAL_FILE.name} (sheet: {ws.title})"))
    return out


def build_courses() -> pd.DataFrame:
    rows, seen = [], set()
    # Classroom Allocation is the most complete/official list; the other two only fill gaps.
    for row in _classroom_allocation() + _ug1_courses() + _finalized_list():
        valid_code = bool(re.fullmatch(r"[A-Z]{2,4}\d{4,5}", row["code"]))
        key = (row["code"], row["slot"]) if valid_code else (row["name"].lower(), row["slot"])
        if valid_code and any(k[0] == row["code"] for k in seen) and "Finalized" in row["source"]:
            continue
        if key in seen:
            continue
        seen.add(key)
        rows.append(row)
    df = pd.DataFrame(rows, columns=COLUMNS)
    df.to_csv(COURSES_CSV, index=False)
    return df


# --------------------------------------------------------------------------- slots
def build_slots() -> pd.DataFrame:
    rows = []
    ws = openpyxl.load_workbook(UG2_FILE, data_only=True)["Slot System"]
    for r in ws.iter_rows(min_row=3, max_row=13, values_only=True):
        time = _s(r[0]).replace("–", "-")
        for d, day in enumerate(DAYS):
            for col in (1 + 2 * d, 2 + 2 * d):  # theory column, lab/tutorial column
                cell = _s(r[col])
                if re.fullmatch(r"[A-Z]{1,2}\d{0,2}", cell):
                    rows.append(dict(slot=cell, day=day, time=time, applies_to="B.Tech. 2nd year onwards / PG",
                                     source=f"{UG2_FILE.name} (sheet: Slot System)"))

    ws = openpyxl.load_workbook(UG1_FILE, data_only=True).worksheets[0]
    for r in ws.iter_rows(min_row=3, max_row=13, values_only=True):
        time = _s(r[0]).replace("–", "-")
        for d, day in enumerate(DAYS):
            cell = _s(r[1 + d])
            m = re.match(r"(A[A-Z]|M\d{1,2})\b(.*)", cell)
            if m:
                rows.append(dict(slot=m.group(1), day=day, time=time + (f" {m.group(2).strip()}" if m.group(2).strip() else ""),
                                 applies_to="B.Tech. 1st year", source=f"{UG1_FILE.name} (sheet: UG 1st Year Slots)"))
    df = pd.DataFrame(rows)
    df.to_csv(SLOTS_CSV, index=False)
    return df


def is_saved_html(path: Path) -> bool:
    """A Google Sheet saved from the browser keeps its .xlsx name but is an HTML page."""
    with open(path, "rb") as f:
        return f.read(64).lstrip().lower().startswith((b"<!doctype html", b"<html"))


def _unescape(v: str) -> str:
    try:
        return json.loads(f'"{v}"')  # the page escapes cell text twice (\\n, \\u0026)
    except json.JSONDecodeError:
        return v


def _gsheet_grid(path: Path) -> list[list[str]]:
    """Cells of a saved Google Sheets page. The page embeds the sheet as one flat, row-major JSON array of
    cells ([] = empty, [{"3": [type, value], ...}] = filled); the column count is the one that puts every
    date in the same column."""
    s = path.read_text(encoding="utf-8").replace('\\"', '"')
    first = re.search(r'"3":\[\d+,"', s)
    cells, _ = json.JSONDecoder().raw_decode(s, s.rfind("[[", 0, first.start()))
    values = []
    for c in cells:
        v = c[0].get("3", [None, ""])[1] if c and isinstance(c[0], dict) else ""
        values.append(_unescape(str(v)).strip() if v is not None else "")
    dated = [i for i, v in enumerate(values) if re.search(r"\d{4}, \w+day", v)]
    ncols = next(n for n in range(2, 60) if len({i % n for i in dated}) == 1)
    return [values[i:i + ncols] for i in range(0, len(values), ncols)]


def sheet_grid(path: Path, sheet: str | None = None) -> list[list[str]]:
    if is_saved_html(path):
        return _gsheet_grid(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [[_s(c) for c in r] for r in ws.iter_rows(values_only=True)]


def _exam_rows(grid: list[list[str]], exam: str, source: str) -> list[dict]:
    """Date-by-session grid (dates in one column, session timings in a header row) -> one row per slot."""
    timings_row = next(r for r in grid if any(re.search(r"[AP]M to", c) for c in r))
    rows = []
    for r in grid:
        dcol = next((i for i, c in enumerate(r) if re.search(r"\d{4}, \w+day", c)), None)
        if dcol is None:
            continue
        date = re.sub(r"([A-Za-z])(\d{4})", r"\1 \2", r[dcol])  # "November2026" -> "November 2026"
        for col, cell in enumerate(r[dcol + 1:], start=dcol + 1):
            for slot in (s.strip() for s in cell.split("+")):
                if re.fullmatch(r"[A-Z]{1,3}\d?", slot):  # skips "Free Slot", "Institute Holiday (...)"
                    rows.append(dict(slot=slot, date=date, time=timings_row[col], exam=exam, source=source))
    return rows


def build_exam_schedule() -> pd.DataFrame:
    df = pd.DataFrame(_exam_rows(sheet_grid(UG2_FILE, "Examination"), "Minor examination",
                                 f"{UG2_FILE.name} (sheet: Examination)"))
    df.to_csv(EXAMS_CSV, index=False)
    return df


def build_major_exam_schedule() -> pd.DataFrame:
    files = sorted(TT_DIR.glob("*Major Exam*.xlsx"))
    rows = _exam_rows(sheet_grid(files[0]), "Major examination", files[0].name) if files else []
    df = pd.DataFrame(rows, columns=["slot", "date", "time", "exam", "source"])
    df.to_csv(MAJOR_EXAMS_CSV, index=False)
    return df


def build_all():
    c, s, e, m = build_courses(), build_slots(), build_exam_schedule(), build_major_exam_schedule()
    print(f"courses.csv: {len(c)} rows | slot_timings.csv: {len(s)} rows | minor_exam_schedule.csv: {len(e)} rows"
          f" | major_exam_schedule.csv: {len(m)} rows")


if __name__ == "__main__":
    build_all()
