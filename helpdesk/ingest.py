"""Load every document under data/raw, split it into chunks, embed with Ollama and store in ChromaDB.

Usage:  python main.py ingest        (or python -m helpdesk.ingest)
"""
import hashlib
import json
import pickle
import re
import subprocess
from collections import Counter
from pathlib import Path

import chromadb
import docx
import fitz  # PyMuPDF
import ollama
import openpyxl
import pandas as pd
from langchain_text_splitters import RecursiveCharacterTextSplitter

from . import config
from .tables import build_all, is_saved_html, sheet_grid

CATEGORY = {"office_orders": "office_order", "course_catalogs": "curriculum",
            "official": "people", "timetables": "timetable"}

# Light redaction of student personal data that appears in e-mail print-outs.
REDACTIONS = [
    (re.compile(r"\b[bmBM]\d{2}[a-zA-Z]{2,3}\d{3,4}@iitj\.ac\.in\b"), "student@iitj.ac.in"),
    (re.compile(r"\b[A-Z][a-z]+ [A-Z][a-z]+ \(([BM]\d{2}[A-Z]{2,3}\d{3,4})\)"), "[Student Name] (B00XX0000)"),
    (re.compile(r"\b[BM]\d{2}[A-Z]{2,3}\d{3,4}\b"), "B00XX0000"),
]


def clean(text: str) -> str:
    text = text.replace("​", "").replace("‌", "").replace(" ", " ")
    for pat, rep in REDACTIONS:
        text = pat.sub(rep, text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    # PDF extraction breaks justified text into one word (or a few) per line. Join soft line breaks so a
    # chunk holds ~30% more text and table rows stay next to their headers; keep paragraph breaks.
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    return text.strip()


# --------------------------------------------------------------------------- loaders
# Each loader returns a list of (page_or_section, text).
def is_scanned(path: Path) -> bool:
    """A PDF without a text layer whose OCR output sits next to it as .txt."""
    if path.suffix.lower() != ".pdf" or not path.with_suffix(".txt").exists():
        return False
    with fitz.open(path) as doc:
        return sum(len(p.get_text().strip()) for p in doc) < 200


def load_pdf(path: Path) -> list[tuple[str, str]]:
    if is_scanned(path):
        return load_txt(path.with_suffix(".txt"))
    with fitz.open(path) as doc:
        return [(str(i + 1), p.get_text()) for i, p in enumerate(doc)]


_WORD = re.compile(r"[A-Za-z]{3,}")


def build_vocabulary(texts: list[str], min_count: int = 3) -> set[str]:
    counts = Counter(w.lower() for t in texts for w in _WORD.findall(t))
    return {w for w, c in counts.items() if c >= min_count}


def drop_ocr_noise(text: str, vocab: set[str], min_ratio: float = 0.5) -> str:
    """The scanned circulars are bilingual; OCR turns the Hindi half into Latin gibberish
    ("ag va ferwar & fear we snare ..."), which dilutes the chunk embeddings. Keep a line only if most of
    its words occur in the vocabulary of the born-digital documents (lines without words are kept)."""
    kept = []
    for line in text.splitlines():
        words = [w.lower() for w in _WORD.findall(line)]
        if not words or sum(w in vocab for w in words) / len(words) >= min_ratio:
            kept.append(line)
    return "\n".join(kept)


def load_txt(path: Path) -> list[tuple[str, str]]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    parts = re.split(r"^## Page (\d+)\s*$", text, flags=re.M)
    if len(parts) == 1:
        return [("1", text)]
    return [(parts[i], parts[i + 1]) for i in range(1, len(parts) - 1, 2)]


def load_docx(path: Path) -> list[tuple[str, str]]:
    d = docx.Document(path)
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return [("1", "\n".join(lines))]


def load_doc(path: Path) -> list[tuple[str, str]]:
    try:  # legacy Word format: antiword ships with Git for Windows
        out = subprocess.run(["antiword", str(path)], capture_output=True, timeout=60)
        return [("1", out.stdout.decode("utf-8", errors="ignore"))]
    except (OSError, subprocess.SubprocessError):
        print(f"  ! antiword not available, skipping {path.name}")
        return []


def load_xlsx(path: Path) -> list[tuple[str, str]]:
    """One section per sheet; every row rendered as 'header: value' pairs so rows survive chunking."""
    if is_saved_html(path):  # Google Sheet saved from the browser: one sheet, rows as "a | b | c"
        rows = [" | ".join(c for c in r if c) for r in sheet_grid(path)]
        return [("sheet 1", "\n".join(r for r in rows if r))]
    wb = openpyxl.load_workbook(path, data_only=True)
    sections = []
    for ws in wb.worksheets:
        lines, header = [], None
        for r in ws.iter_rows(values_only=True):
            cells = [re.sub(r"\s+", " ", str(c)).strip() if c is not None else "" for c in r]
            if not any(cells):
                continue
            filled = [c for c in cells if c]
            if header is None and len(filled) >= 3 and any(k in " ".join(filled).lower() for k in ("code", "slot", "name", "subject")):
                header = cells
                continue
            if header:
                pairs = [f"{header[i] or 'col'+str(i)}: {c}" for i, c in enumerate(cells) if c and i < len(header)]
                lines.append("; ".join(pairs))
            else:
                lines.append(" | ".join(filled))
        if lines:
            sections.append((f"sheet {ws.title}", "\n".join(lines)))
    return sections


def load_csv(path: Path) -> list[tuple[str, str]]:
    df = pd.read_csv(path).fillna("")
    lines = ["; ".join(f"{k}: {v}" for k, v in row.items() if v) for row in df.to_dict("records")]
    return [("rows", "\n".join(lines))]


LOADERS = {".pdf": load_pdf, ".txt": load_txt, ".docx": load_docx, ".doc": load_doc,
           ".xlsx": load_xlsx, ".csv": load_csv}


def discover_files() -> list[Path]:
    files = []
    for p in sorted(config.RAW_DIR.rglob("*")):
        if not p.is_file() or p.suffix.lower() not in LOADERS:
            continue
        if any(x in p.name for x in config.EXCLUDE_PATTERNS):
            continue
        # .txt files are OCR companions of scanned PDFs; they are read through load_pdf
        if p.suffix.lower() == ".txt" and p.with_suffix(".pdf").exists():
            continue
        files.append(p)
    return files


# --------------------------------------------------------------------------- chunking
def make_chunks(files: list[Path]) -> list[dict]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=config.CHUNK_SIZE, chunk_overlap=config.CHUNK_OVERLAP,
                                              separators=["\n\n", "\n", ". ", " ", ""])
    loaded = [(path, LOADERS[path.suffix.lower()](path), is_scanned(path)) for path in files]
    vocab = build_vocabulary([t for p, secs, scanned in loaded if not scanned and p.suffix.lower() in (".pdf", ".docx")
                              for _, t in secs])
    chunks, seen_docs = [], set()
    for path, sections, scanned in loaded:
        if scanned:
            sections = [(page, drop_ocr_noise(text, vocab)) for page, text in sections]
        full = clean("\n".join(t for _, t in sections))
        digest = hashlib.md5(full.encode()).hexdigest()
        if not full or digest in seen_docs:  # skip empty and byte-identical duplicates
            print(f"  - skipped (empty/duplicate): {path.name}")
            continue
        seen_docs.add(digest)
        category = CATEGORY.get(path.parent.name, "other")
        n = 0
        for page, text in sections:
            for piece in splitter.split_text(clean(text)):
                if len(piece) < 40:
                    continue
                chunks.append(dict(id=f"{digest[:10]}-{len(chunks)}", text=piece,
                                   meta=dict(source=path.name, page=page, category=category)))
                n += 1
        print(f"  + {n:4d} chunks  {path.name}")
    return chunks


CACHE_FILE = config.INDEX_DIR / "embedding_cache.pkl"


def _key(text: str) -> str:
    return hashlib.sha1(f"{config.EMBED_MODEL}\n{text}".encode()).hexdigest()


def load_cache() -> dict[str, list[float]]:
    """Embeddings keyed by chunk text, so re-indexing only embeds new or changed chunks
    (a full embedding pass takes over an hour on a 2-core CPU)."""
    if CACHE_FILE.exists():
        return pickle.loads(CACHE_FILE.read_bytes())
    cache = {}
    try:  # seed from an index built before the cache existed
        col = chromadb.PersistentClient(path=str(config.INDEX_DIR / "chroma")).get_collection(config.COLLECTION)
        data = col.get(include=["documents", "embeddings"])
        cache = {_key(d): list(map(float, e)) for d, e in zip(data["documents"], data["embeddings"])}
    except Exception:
        pass
    return cache


def embed(texts: list[str], kind: str = "document", batch: int = 32) -> list[list[float]]:
    # nomic-embed-text was trained with task prefixes
    prefix = "search_document: " if kind == "document" else "search_query: "
    vectors = []
    for i in range(0, len(texts), batch):
        resp = ollama.embed(model=config.EMBED_MODEL, input=[prefix + t for t in texts[i:i + batch]])
        vectors.extend(resp["embeddings"])
        if kind == "document":
            print(f"\r  embedded {min(i + batch, len(texts))}/{len(texts)}", end="", flush=True)
    if kind == "document" and texts:
        print()
    return vectors


def build_index():
    print("1) Building structured tables from timetable spreadsheets ...")
    build_all()

    print("2) Loading and chunking documents ...")
    files = discover_files()
    chunks = make_chunks(files)
    print(f"   {len(files)} files -> {len(chunks)} chunks")

    print(f"3) Embedding with {config.EMBED_MODEL} and writing ChromaDB ...")
    cache = load_cache()
    todo = list(dict.fromkeys(c["text"] for c in chunks if _key(c["text"]) not in cache))
    print(f"   {len(chunks) - len(todo)} embeddings cached, {len(todo)} to compute")
    cache.update({_key(t): v for t, v in zip(todo, embed(todo))})
    CACHE_FILE.write_bytes(pickle.dumps(cache))
    vectors = [cache[_key(c["text"])] for c in chunks]

    # delete_collection rather than removing the folder: on Windows the sqlite file stays locked by the
    # client that seeded the cache, and a silently failed rmtree leaves a half-deleted index.
    client = chromadb.PersistentClient(path=str(config.INDEX_DIR / "chroma"))
    if config.COLLECTION in [c.name for c in client.list_collections()]:
        client.delete_collection(config.COLLECTION)
    col = client.create_collection(config.COLLECTION, metadata={"hnsw:space": "cosine"})
    for i in range(0, len(chunks), 500):
        part = chunks[i:i + 500]
        col.add(ids=[c["id"] for c in part], documents=[c["text"] for c in part],
                metadatas=[c["meta"] for c in part], embeddings=vectors[i:i + 500])

    manifest = sorted({c["meta"]["source"] for c in chunks})
    (config.INDEX_DIR / "manifest.json").write_text(json.dumps(manifest, indent=1), encoding="utf-8")
    print(f"Done: {col.count()} chunks from {len(manifest)} documents.")


if __name__ == "__main__":
    build_index()
